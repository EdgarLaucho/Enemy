from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\jdavi\Documents\Unreal Projects\Enemy")
OUTPUT = ROOT / "output" / "docx" / "Carta_presentacion_Juan_David_Iglesias_Indra_Group.docx"

NAVY = RGBColor(48, 82, 121)
CHARCOAL = RGBColor(38, 38, 38)
MUTED = RGBColor(92, 92, 92)


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_text(paragraph, text, size=10.6, bold=False, color=CHARCOAL, italic=False):
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return run


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.6)
    normal.font.color.rgb = CHARCOAL
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.1

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(19)
    title.font.bold = True
    title.font.color.rgb = CHARCOAL
    title.paragraph_format.space_after = Pt(11)
    title_p_pr = title.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    if "Contact line" not in [style.name for style in doc.styles]:
        contact_style = doc.styles.add_style("Contact line", WD_STYLE_TYPE.PARAGRAPH)
    else:
        contact_style = doc.styles["Contact line"]
    contact_style.font.name = "Arial"
    contact_style.font.size = Pt(9)
    contact_style.font.color.rgb = MUTED
    contact_style.paragraph_format.space_after = Pt(2)

    header = doc.add_table(rows=1, cols=2)
    header.autofit = False
    header.columns[0].width = Inches(4.6)
    header.columns[1].width = Inches(2.0)
    header.rows[0].cells[0].width = Inches(4.6)
    header.rows[0].cells[1].width = Inches(2.0)
    header.style = "Table Grid"
    for cell in header.rows[0].cells:
        set_cell_margins(cell, top=80, start=110, bottom=85, end=110)
        shade_cell(cell, "355F8C")
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = tc_pr.find(qn("w:tcBorders"))
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tc_pr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            tag = borders.find(qn(f"w:{edge}"))
            if tag is None:
                tag = OxmlElement(f"w:{edge}")
                borders.append(tag)
            tag.set(qn("w:val"), "nil")

    p = header.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_text(p, "JUAN DAVID\n", size=11.5, color=RGBColor(255, 255, 255))
    add_text(p, "IGLESIAS MARTÍNEZ", size=16.5, bold=True, color=RGBColor(255, 255, 255))

    p = header.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    add_text(p, "DESARROLLADOR\nJAVA JUNIOR", size=10.5, bold=True, color=RGBColor(255, 255, 255))

    p = doc.add_paragraph(style="Contact line")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(
        p,
        "Leganés, Madrid  •  608 893 746  •  jdavidiglesias1991@gmail.com  •  github.com/sekiro91",
        size=8.9,
        color=MUTED,
    )
    p.paragraph_format.space_after = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(8)
    add_text(p, "Madrid, 16 de septiembre de 2026", size=10, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    add_text(p, "A la atención del equipo de Selección", size=10.4, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    add_text(p, "Indra Group", size=10.4, bold=True, color=NAVY)

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.keep_with_next = True
    add_text(p, "Carta de presentación", size=19, bold=True, color=CHARCOAL)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    add_text(p, "Asunto: ", size=10.6, bold=True, color=CHARCOAL)
    add_text(p, "Candidatura para Desarrollador Java Junior", size=10.6, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    add_text(p, "Estimado equipo de Selección:")

    paragraphs = [
        (
            "Me gustaría presentar mi candidatura para incorporarme a Indra Group como Desarrollador Java Junior. "
            "He finalizado el Ciclo Formativo de Grado Superior en Desarrollo de Aplicaciones Web (DAW), "
            "incluidas las prácticas profesionales (FCT) en empresa, y recibiré el título oficial en enero de 2027. "
            "Cuento con conocimientos de Java y programación orientada a objetos, además de una base técnica en "
            "desarrollo web full-stack."
        ),
        (
            "Desde mayo de 2026 participo en el desarrollo de Hookit, una plataforma SaaS en producción construida "
            "con Next.js y React. He entregado más de 200 commits propios mediante pull requests revisadas por el "
            "equipo y he desarrollado funcionalidades completas: interfaces, páginas dinámicas, formularios de "
            "captación, control de acceso y automatizaciones de correo. Esta experiencia me ha permitido trabajar "
            "con Git, revisión de código y criterios de calidad en un producto real."
        ),
        (
            "Aporto también siete años de experiencia en atención al cliente y venta consultiva, entre ellos como "
            "teleoperador bilingüe en Movistar Prosegur Alarmas, donde resolvía incidencias técnicas, reclamaciones "
            "y gestiones con usuarios. Esa trayectoria ha reforzado mi capacidad para analizar problemas, comunicarme "
            "con claridad, asumir responsabilidades y mantener el foco en las necesidades del cliente. Además, tengo "
            "un nivel B2 de inglés tras haber trabajado dos años en Dublín."
        ),
        (
            "Me interesa Indra Group por la oportunidad de crecer en equipos tecnológicos, participar en proyectos "
            "de alcance y consolidar mi especialización en Java. Puedo aportar motivación, experiencia práctica en "
            "desarrollo, capacidad de aprendizaje y una orientación al usuario poco habitual en un perfil junior."
        ),
        (
            "Quedo a su disposición para ampliar cualquier información en una entrevista. Muchas gracias por valorar "
            "mi candidatura."
        ),
    ]

    for idx, text in enumerate(paragraphs):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(7 if idx < len(paragraphs) - 1 else 10)
        p.paragraph_format.widow_control = True
        add_text(p, text)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    add_text(p, "Atentamente,")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    add_text(p, "Juan David Iglesias Martínez", size=11, bold=True, color=NAVY)
    p = doc.add_paragraph(style="Contact line")
    add_text(p, "608 893 746  •  jdavidiglesias1991@gmail.com", size=9, color=MUTED)

    doc.core_properties.title = "Carta de presentación — Juan David Iglesias Martínez — Indra Group"
    doc.core_properties.subject = "Candidatura para Desarrollador Java Junior"
    doc.core_properties.author = "Juan David Iglesias Martínez"
    doc.core_properties.keywords = "Indra Group, Java, desarrollador junior, DAW"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
